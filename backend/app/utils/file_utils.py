import os
from pathlib import Path

from fastapi import UploadFile

from app.config import settings


ALLOWED_EXTENSIONS = settings.allowed_extensions


def safe_filename(name: str) -> str:
    base = os.path.basename(name).strip()
    safe = "".join(ch for ch in base if ch.isalnum() or ch in (".", "_", "-"))
    return safe or "document"


def ensure_upload_directory() -> str:
    upload_dir = Path(settings.upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)
    return str(upload_dir)


def validate_upload(file: UploadFile) -> None:
    if not file.filename:
        raise ValueError("Filename is required.")
    extension = os.path.splitext(file.filename)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Allowed: PDF, DOCX, TXT")

    expected_content_types = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
    }
    if file.content_type and file.content_type != "application/octet-stream":
        if file.content_type not in settings.allowed_content_types or file.content_type != expected_content_types[extension]:
            raise ValueError("File extension and content type do not match.")

    size = file.size
    if size is None:
        position = file.file.tell()
        file.file.seek(0, os.SEEK_END)
        size = file.file.tell()
        file.file.seek(position)
    if size == 0:
        raise ValueError("Uploaded file is empty.")
    if size > settings.max_upload_size:
        raise ValueError("File too large.")


def extract_pages_from_file(file_path: str) -> list[tuple[int, str]]:
    extension = Path(file_path).suffix.lower()
    if extension == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
            return [(1, handle.read())]
    if extension == ".pdf":
        import fitz

        with fitz.open(file_path) as doc:
            return [(page_number, page.get_text()) for page_number, page in enumerate(doc, start=1)]
    if extension == ".docx":
        import docx

        document = docx.Document(file_path)
        content = [paragraph.text for paragraph in document.paragraphs]
        content.extend("\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return [(1, "\n".join(content))]
    raise ValueError(f"Unsupported document type: {extension}")


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 180) -> list[str]:
    if not text.strip():
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end]
        chunks.append(chunk.strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]
