import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from google.genai.errors import ServerError
from io import BytesIO
from types import SimpleNamespace
from starlette.datastructures import Headers, UploadFile

from app.config import settings
from app.database import SessionLocal
from app.main import app
from app.models.document import Document, DocumentChunk
from app.models.user import User
from app.services.document_service import DocumentService
from app.services.gemini_service import GeminiService
from app.services.vector_service import VectorService
from app.utils.file_utils import chunk_text, extract_pages_from_file, validate_upload
from tests.conftest import reset_db

client = TestClient(app)


def test_pdf_extraction_preserves_each_page(tmp_path):
    import fitz

    pdf_path = tmp_path / "two-pages.pdf"
    with fitz.open() as pdf:
        first_page = pdf.new_page()
        first_page.insert_text((72, 72), "First page content")
        second_page = pdf.new_page()
        second_page.insert_text((72, 72), "Second page content")
        pdf.save(pdf_path)

    extracted_pages = extract_pages_from_file(str(pdf_path))

    assert extracted_pages == [(1, "First page content\n"), (2, "Second page content\n")]


def test_text_extraction_and_cleaning(tmp_path):
    text_path = tmp_path / "notes.txt"
    text_path.write_text("  First line\n\nSecond   line  ", encoding="utf-8")

    pages = extract_pages_from_file(str(text_path))

    assert pages == [(1, "  First line\n\nSecond   line  ")]
    assert " ".join(pages[0][1].split()) == "First line Second line"


def test_chunking_preserves_overlap_between_chunks():
    chunks = chunk_text("ABCDEFGHIJKLM", chunk_size=8, overlap=3)

    assert chunks == ["ABCDEFGH", "FGHIJKLM"]


def test_docx_extraction_includes_paragraphs_and_tables(tmp_path):
    from docx import Document as DocxDocument

    docx_path = tmp_path / "notes.docx"
    document = DocxDocument()
    document.add_paragraph("Document paragraph")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = "Table cell one"
    cells[1].text = "Table cell two"
    document.save(docx_path)

    pages = extract_pages_from_file(str(docx_path))

    assert pages == [(1, "Document paragraph\nTable cell one\tTable cell two")]


def test_upload_validation_rejects_mismatched_content_type():
    from io import BytesIO
    from starlette.datastructures import Headers, UploadFile

    upload = UploadFile(
        file=BytesIO(b"not a PDF"),
        filename="notes.pdf",
        headers=Headers({"content-type": "text/plain"}),
    )

    try:
        validate_upload(upload)
    except ValueError as exc:
        assert str(exc) == "File extension and content type do not match."
    else:
        raise AssertionError("Mismatched upload content type should be rejected.")


def test_upload_validation_checks_size_when_upload_size_is_unknown(monkeypatch):
    from io import BytesIO
    from starlette.datastructures import Headers, UploadFile

    monkeypatch.setattr(settings, "max_upload_size", 4)
    upload = UploadFile(
        file=BytesIO(b"12345"),
        filename="notes.txt",
        headers=Headers({"content-type": "text/plain"}),
    )

    try:
        validate_upload(upload)
    except ValueError as exc:
        assert str(exc) == "File too large."
    else:
        raise AssertionError("Oversized upload should be rejected even when size metadata is absent.")


def test_summary_generation_reports_gemini_quota_without_fabricating_summary():
    reset_db()
    db = SessionLocal()
    user = User(email="summary-quota@example.com", name="Summary", password_hash="unused")
    db.add(user)
    db.commit()
    document = Document(
        user_id=user.id,
        filename="summary.txt",
        storage_name="summary.txt",
        content_type="text/plain",
        file_size=10,
        status="READY",
    )
    db.add(document)
    db.commit()
    db.add(DocumentChunk(document_id=document.id, chunk_index=0, content="Persisted summary context", page=1, source_filename="summary.txt", chunk_metadata={}))
    db.commit()
    service = DocumentService(db)

    def raise_quota_error(prompt, system_instruction=None):
        raise ServerError(429, {"error": {"status": "RESOURCE_EXHAUSTED", "message": "quota exhausted"}})

    service.gemini.generate_text = raise_quota_error

    with pytest.raises(HTTPException) as raised:
        service.generate_summary(user.id, document.id)

    assert raised.value.status_code == 429
    assert "No summary was generated" in raised.value.detail
    assert document.summary is None
    db.close()


def test_document_upload_and_ownership(monkeypatch, tmp_path):
    reset_db()
    monkeypatch.setattr(settings, "chroma_persist_directory", str(tmp_path / "vectors"))
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    monkeypatch.setattr(GeminiService, "embed_documents", lambda self, texts: [[0.1, 0.2, 0.3] for _ in texts])
    first = client.post("/api/auth/register", json={"name": "User One", "email": "one@example.com", "password": "abc123"})
    second = client.post("/api/auth/register", json={"name": "User Two", "email": "two@example.com", "password": "abc123"})

    token_one = first.json()["access_token"]
    token_two = second.json()["access_token"]

    content = b"This is a sample study document for testing." 
    response = client.post(
        "/api/documents/upload",
        files={"file": ("sample.txt", content, "text/plain")},
        headers={"Authorization": f"Bearer {token_one}"},
    )
    assert response.status_code == 200
    document_id = response.json()["id"]

    owned_document = client.get(f"/api/documents/{document_id}", headers={"Authorization": f"Bearer {token_one}"})
    assert owned_document.status_code == 200

    forbidden = client.get(f"/api/documents/{document_id}", headers={"Authorization": f"Bearer {token_two}"})
    assert forbidden.status_code == 404


def test_document_processing_preserves_chunks_and_metadata_across_embedding_batches(tmp_path, monkeypatch):
    reset_db()
    monkeypatch.setattr(settings, "chroma_persist_directory", str(tmp_path / "vectors"))
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    db = SessionLocal()
    user = User(email="batch-upload@example.com", name="Batch Upload", password_hash="unused")
    db.add(user)
    db.commit()

    class OrderedEmbeddingModels:
        def __init__(self):
            self.calls = []
            self.next_index = 0

        def embed_content(self, model, contents):
            self.calls.append(list(contents))
            count = len(contents)
            vectors = [
                SimpleNamespace(values=[float(index), 1.0])
                for index in range(self.next_index, self.next_index + count)
            ]
            self.next_index += count
            return SimpleNamespace(embeddings=vectors)

    models = OrderedEmbeddingModels()
    service = DocumentService(db)
    service.gemini.client = SimpleNamespace(models=models)
    upload = UploadFile(
        file=BytesIO(b"A" * 164181),
        filename="long.txt",
        headers=Headers({"content-type": "text/plain"}),
    )

    document = service.process_document(user.id, upload)
    chunks = db.query(DocumentChunk).filter_by(document_id=document.id).order_by(DocumentChunk.chunk_index).all()
    stored = service.vector_service.get_collection(user.id).get(
        ids=[f"{document.id}_{index}" for index in range(len(chunks))],
        include=["embeddings", "metadatas"],
    )

    assert document.status == "READY"
    assert len(chunks) == 201
    assert [len(batch) for batch in models.calls] == [100, 100, 1]
    assert len(stored["ids"]) == len(chunks)
    assert [int(vector[0]) for vector in stored["embeddings"]] == list(range(len(chunks)))
    assert [int(metadata["chunk_id"]) for metadata in stored["metadatas"]] == list(range(len(chunks)))
    assert all(metadata["user_id"] == str(user.id) for metadata in stored["metadatas"])
    assert all(metadata["document_id"] == str(document.id) for metadata in stored["metadatas"])
    assert all(metadata["page"] == 1 for metadata in stored["metadatas"])
    db.close()


def test_later_embedding_batch_failure_leaves_document_failed_without_partial_vectors(tmp_path, monkeypatch):
    reset_db()
    monkeypatch.setattr(settings, "chroma_persist_directory", str(tmp_path / "vectors"))
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    db = SessionLocal()
    user = User(email="failed-batch@example.com", name="Failed Batch", password_hash="unused")
    db.add(user)
    db.commit()

    class FailingSecondBatchModels:
        def __init__(self):
            self.batch_sizes = []
            self.embedding_index = 0

        def embed_content(self, model, contents):
            self.batch_sizes.append(len(contents))
            if len(self.batch_sizes) == 2:
                raise ServerError(400, {"error": {"status": "INVALID_ARGUMENT", "message": "diagnostic second-batch failure"}})
            vectors = [
                SimpleNamespace(values=[float(index), 1.0])
                for index in range(self.embedding_index, self.embedding_index + len(contents))
            ]
            self.embedding_index += len(contents)
            return SimpleNamespace(embeddings=vectors)

    models = FailingSecondBatchModels()
    service = DocumentService(db)
    service.gemini.client = SimpleNamespace(models=models)
    upload = UploadFile(
        file=BytesIO(b"A" * 164181),
        filename="later-batch-failure.txt",
        headers=Headers({"content-type": "text/plain"}),
    )

    with pytest.raises(HTTPException) as raised:
        service.process_document(user.id, upload)

    document = db.query(Document).filter_by(user_id=user.id).order_by(Document.id.desc()).first()
    stored = service.vector_service.get_collection(user.id).get(where={"document_id": str(document.id)})
    chunk_count = db.query(DocumentChunk).filter_by(document_id=document.id).count()

    assert models.batch_sizes == [100, 100]
    assert raised.value.status_code == 400
    assert document.status == "FAILED"
    assert "diagnostic second-batch failure" in document.error
    assert chunk_count == 0
    assert stored["ids"] == []
    db.close()


def test_invalid_document_upload():
    reset_db()
    user = client.post("/api/auth/register", json={"name": "Alice", "email": "alice@example.com", "password": "password123"})
    token = user.json()["access_token"]
    response = client.post(
        "/api/documents/upload",
        files={"file": ("bad.exe", b"not a valid file", "application/x-msdownload")},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 400


def test_vector_service_inserts_and_retrieves_from_chromadb(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "chroma_persist_directory", str(tmp_path))
    service = VectorService()
    assert service.client is not None

    service.add_documents(
        user_id=42,
        document_id=7,
        chunks=["alpha", "beta"],
        embeddings=[[0.1, 0.2], [0.3, 0.4]],
        metadata=[{"user_id": "42", "document_id": "7"}, {"user_id": "42", "document_id": "7"}],
    )
    service.add_documents(
        user_id=42,
        document_id=8,
        chunks=["different document"],
        embeddings=[[0.5, 0.6]],
        metadata=[{"user_id": "42", "document_id": "8"}],
    )
    service.add_documents(
        user_id=43,
        document_id=7,
        chunks=["different user"],
        embeddings=[[0.1, 0.2]],
        metadata=[{"user_id": "43", "document_id": "7"}],
    )

    results = service.search(user_id=42, document_id=7, query_embedding=[0.1, 0.2], limit=2)
    assert results["ids"][0][0] == "7_0"
    assert results["documents"][0][0] == "alpha"
    assert set(results["documents"][0]) == {"alpha", "beta"}
    other_user_results = service.search(user_id=43, document_id=7, query_embedding=[0.1, 0.2], limit=2)
    assert other_user_results["documents"][0] == ["different user"]
